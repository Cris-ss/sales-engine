"""Exportação CSV/XLSX de TODOS os leads do universo filtrado (não só a página),
com a mesma consulta-base, filtros e ordenação da lista.

- CNPJ, CEP e telefones saem formatados com pontuação e como TEXTO (Excel não
  remove zeros à esquerda de "00.000.000/0001-91"); no XLSX a célula é tipada
  como string com formato "@".
- Células de texto que começariam com = + - @ (ou tab/CR) recebem um apóstrofo
  na frente para não serem interpretadas como fórmula (CSV/Excel injection).
- CSV: ponto e vírgula + UTF-8 com BOM (abre direto no Excel em português).
- Não exporta erros internos de envio nem segredos.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import datetime
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from api.adapters.whatsapp_adapter import telefone_apto
from api.dependencies import FiltrosLeads
from api.queries.lead_projection import LeadProjection
from api.services.lead_service import montar_condicoes

CABECALHO = [
    "ID", "CNPJ", "Nome", "Nicho", "Município", "UF", "CEP", "Logradouro", "Número", "Bairro",
    "Score", "Faixa do score", "Coluna do funil", "Estágio", "Estágio derivado (sem histórico)",
    "Email confirmado", "Email", "Formulário confirmado", "Formulário (URL)",
    "Telefone 1", "Telefone 2", "Telefone apto a link WhatsApp (não verificado)",
    "Site", "Status do email", "Emails enviados", "Último email enviado em", "Motivo (fora do fluxo)",
]
COLUNAS_TEXTO = {"CNPJ", "CEP", "Telefone 1", "Telefone 2", "ID"}
_PERIGOSOS = ("=", "+", "-", "@", "\t", "\r")


def formatar_cnpj(v: Optional[str]) -> str:
    d = re.sub(r"\D", "", v or "")
    if len(d) != 14:
        return v or ""
    return f"{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}"


def formatar_cep(v: Optional[str]) -> str:
    d = re.sub(r"\D", "", v or "")
    return f"{d[:5]}-{d[5:]}" if len(d) == 8 else (v or "")


def formatar_telefone(v: Optional[str]) -> str:
    d = re.sub(r"\D", "", v or "")
    if len(d) == 10:
        return f"({d[:2]}) {d[2:6]}-{d[6:]}"
    if len(d) == 11:
        return f"({d[:2]}) {d[2:7]}-{d[7:]}"
    return v or ""


def neutralizar(v: Any) -> Any:
    if isinstance(v, str) and v.startswith(_PERIGOSOS):
        return "'" + v
    return v


def _sim_nao(b: Any) -> str:
    return "sim" if b else "não"


def _data(dt: Optional[datetime]) -> str:
    return dt.astimezone().strftime("%d/%m/%Y %H:%M") if dt else ""


def linhas(session: Session, f: FiltrosLeads) -> list[list[Any]]:
    p = LeadProjection()
    conds = montar_condicoes(session, p, f)
    rows = session.execute(
        p.select(p.colunas_detalhe()).where(*conds).order_by(*p.ordenacao(f.sort, f.order))
    ).mappings().all()

    saida = []
    for r in rows:
        email_ok, form_ok = bool(r["email_confirmado"]), bool(r["formulario_confirmado"])
        saida.append([
            str(r["id"]), formatar_cnpj(r["cnpj"]), r["nome"], r["nicho_nome"], r["municipio"] or "", r["uf"] or "",
            formatar_cep(r["cep"]), r["logradouro"] or "", r["numero"] or "", r["bairro"] or "",
            r["score"] if r["score"] is not None else "", r["faixa_score"], r["coluna"], r["estagio"],
            _sim_nao(r["estagio_derivado"]),
            _sim_nao(email_ok), (r["email_final"] or "") if email_ok else "",
            _sim_nao(form_ok), (r["formulario_contato_url"] or "") if form_ok else "",
            formatar_telefone(r["telefone1"]), formatar_telefone(r["telefone2"]),
            _sim_nao(telefone_apto(r["telefone1"], r["telefone2"])),
            r["site_url"] or "", r["email_status"], int(r["qtd_emails_enviados"] or 0),
            _data(r["email_enviado_em"]), r["motivo_reprovacao"] or "",
        ])
    return saida


def _celula(valor: Any) -> Any:
    return neutralizar(valor)


def gerar_csv(dados: list[list[Any]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    w.writerow(CABECALHO)
    for linha in dados:
        w.writerow([_celula(v) for v in linha])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def gerar_xlsx(dados: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(CABECALHO)
    for c in ws[1]:
        c.font = Font(bold=True)

    idx_texto = {i for i, nome in enumerate(CABECALHO) if nome in COLUNAS_TEXTO}
    for linha in dados:
        ws.append([None] * len(linha))
        n = ws.max_row
        for i, v in enumerate(linha):
            cel = ws.cell(row=n, column=i + 1)
            v = _celula(v)
            if i in idx_texto or isinstance(v, str):
                cel.value = v
                cel.data_type = "s"            # força string: sem coerção para número/fórmula
                if i in idx_texto:
                    cel.number_format = "@"
            else:
                cel.value = v

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CABECALHO))}{max(ws.max_row, 1)}"
    for i, nome in enumerate(CABECALHO, start=1):
        maior = max([len(nome)] + [len(str(l[i - 1])) for l in dados[:200]] or [len(nome)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, maior + 2), 48)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
