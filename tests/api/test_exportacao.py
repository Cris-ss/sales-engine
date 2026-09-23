"""Exportação: bate com a consulta filtrada
completa, mantém zeros à esquerda, neutraliza fórmulas, não vaza erro interno."""

from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

from api.services.export_service import CABECALHO


def _csv(client, qs=""):
    r = client.get(f"/api/v1/exportacoes/leads?formato=csv{qs}")
    assert r.status_code == 200
    assert r.content.startswith("﻿".encode("utf-8"))            # BOM p/ Excel
    linhas = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
    return r, linhas[0], linhas[1:]


def test_exporta_todos_os_filtrados_e_nao_so_a_pagina(client, fab, db):
    for i in range(130):
        e = fab.empresa(f"Emp{i:03d}", uf="PR" if i < 120 else "SP")
        fab.score(e, 55)
    db.commit()
    lista = client.get("/api/v1/leads?page_size=100&uf=PR").json()
    assert lista["total"] == 120 and len(lista["itens"]) == 100      # a página não tem tudo
    r, cab, linhas = _csv(client, "&uf=PR")
    assert cab == CABECALHO and len(linhas) == 120 == int(r.headers["X-Total-Linhas"])
    ids_export = [l[0] for l in linhas]
    ids_lista = [i["id"] for i in client.get("/api/v1/leads?page_size=100&uf=PR").json()["itens"]]
    assert [int(i) for i in ids_export[:100]] == ids_lista           # mesma ordenação da consulta-base
    assert all(l[CABECALHO.index("UF")] == "PR" for l in linhas)


def test_identificadores_mantem_zeros_e_pontuacao_como_texto(client, fab, db):
    fab.empresa("Zero", tel1="1199180999", tem_pelo_menos_um_canal=True, email_final="z@z.com")
    db.commit()
    _, cab, linhas = _csv(client)
    l = dict(zip(cab, linhas[0]))
    assert l["CNPJ"] == "00.000.000/0000-01"          # zeros à esquerda preservados
    assert l["CEP"] == "12345-678"
    assert l["Telefone 1"] == "(11) 9918-0999"        # formato original da Receita (8 dígitos), sem inventar o 9
    assert l["Telefone apto a link WhatsApp (não verificado)"] == "sim"


def test_xlsx_cabecalho_filtro_congelamento_e_texto(client, fab, db):
    fab.empresa("Alfa", email_final="a@a.com", tem_pelo_menos_um_canal=True)
    fab.empresa("Beta")
    db.commit()
    r = client.get("/api/v1/exportacoes/leads?formato=xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Leads"]
    assert [c.value for c in ws[1]] == CABECALHO
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1:")
    assert ws.max_row == 3
    cnpj = ws.cell(row=2, column=CABECALHO.index("CNPJ") + 1)
    assert cnpj.data_type == "s" and cnpj.value.startswith("00.")


def test_neutraliza_formulas_em_csv_e_xlsx(client, fab, db):
    fab.empresa("=HYPERLINK(\"http://x\";\"clique\")")
    fab.empresa("+SOMA(1;1)")
    db.commit()
    _, cab, linhas = _csv(client)
    nomes = {l[cab.index("Nome")] for l in linhas}
    assert all(n.startswith("'") for n in nomes)
    ws = load_workbook(io.BytesIO(client.get("/api/v1/exportacoes/leads?formato=xlsx").content))["Leads"]
    for row in range(2, 4):
        cel = ws.cell(row=row, column=CABECALHO.index("Nome") + 1)
        assert cel.data_type != "f" and cel.value.startswith("'")


def test_nao_exporta_erros_internos_nem_email_nao_confirmado(client, fab, db):
    e = fab.empresa("Erro", email_final="oculto@x.com", tem_pelo_menos_um_canal=False, motivo_reprovacao="site_nao_confirmado_marca_ausente")
    fab.contato(e, status="pendente", erro="SMTPAuthenticationError: segredo interno")
    db.commit()
    r, cab, linhas = _csv(client)
    corpo = r.content.decode("utf-8-sig")
    assert "segredo interno" not in corpo and "oculto@x.com" not in corpo
    l = dict(zip(cab, linhas[0]))
    assert l["Email confirmado"] == "não" and l["Motivo (fora do fluxo)"] == "site_nao_confirmado_marca_ausente"


def test_formato_invalido_e_filtros_invalidos(client):
    assert client.get("/api/v1/exportacoes/leads?formato=pdf").status_code == 422
    assert client.get("/api/v1/exportacoes/leads?canal=x").status_code == 422
